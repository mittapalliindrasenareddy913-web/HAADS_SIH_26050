import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'SIH2026_template_working.pptx')
print(f'Total Slides: {len(prs.slides)}')
print(f'Slide Width: {prs.slide_width.inches:.3f} in, Height: {prs.slide_height.inches:.3f} in')

for idx, slide in enumerate(prs.slides):
    print(f'\n================ SLIDE {idx+1} ================')
    print(f'Layout Name: {slide.slide_layout.name}')
    for s_idx, shape in enumerate(slide.shapes):
        name = shape.name
        left, top, w, h = shape.left.inches, shape.top.inches, shape.width.inches, shape.height.inches
        has_tf = shape.has_text_frame
        text = shape.text.strip().replace('\n', ' ') if has_tf else ''
        print(f' Shape {s_idx+1}: [{name}] Pos=({left:.2f}, {top:.2f}, {w:.2f}, {h:.2f})')
        if text:
            print(f'    Text: "{text[:140]}"')
