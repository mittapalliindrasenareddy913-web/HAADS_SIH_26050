import sys
import pptx

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation(r'SIH2026_template_working.pptx')

for idx, slide in enumerate(prs.slides):
    print(f'================ SLIDE {idx+1} ================')
    for s_idx, shape in enumerate(slide.shapes):
        print(f'Shape {s_idx+1}: ID={shape.shape_id}, Name="{shape.name}", Type={shape.shape_type}')
        if shape.has_text_frame:
            for p_idx, p in enumerate(shape.text_frame.paragraphs):
                runs_text = "".join(r.text for r in p.runs)
                print(f'   P{p_idx+1}: "{p.text}" | Runs: "{runs_text}"')
